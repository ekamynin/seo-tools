import io
import zipfile
from unittest.mock import Mock, patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

from content_publisher.models import Block, DocumentResult, InlinePart
from content_publisher.google_docs import (
    GoogleDocsError,
    download_google_doc,
    extract_google_doc_id,
    extract_google_doc_links,
)
from content_publisher.parser import parse_docx
from content_publisher.renderer import build_docx, build_docx_zip, render_html, validate_html
from content_publisher.text_parser import parse_plain_text, parse_rich_text


def _docx_bytes(build) -> bytes:
    document = Document()
    build(document)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        RELATIONSHIP_TYPE.HYPERLINK,
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def test_heading_levels_are_normalized_from_smallest_word_level():
    def build(document):
        document.add_heading("Основний розділ", level=2)
        document.add_paragraph("Звичайний текст статті, який має залишитися абзацом.")
        document.add_heading("Підрозділ", level=3)

    result = parse_docx(_docx_bytes(build), "article.docx")

    assert [block.role for block in result.blocks] == ["h2", "p", "h3"]


def test_bold_inside_paragraph_is_not_rendered():
    def build(document):
        paragraph = document.add_paragraph("Купити ")
        run = paragraph.add_run("ключове слово")
        run.bold = True
        paragraph.add_run(" в Україні.")

    result = parse_docx(_docx_bytes(build), "article.docx")
    fragment = render_html(result.blocks)

    assert fragment == "<p>Купити ключове слово в Україні.</p>"
    assert "<strong" not in fragment
    assert "<b" not in fragment


def test_fully_bold_short_paragraph_is_heading_candidate():
    def build(document):
        paragraph = document.add_paragraph()
        run = paragraph.add_run("Історія бренду")
        run.bold = True
        document.add_paragraph(
            "Це достатньо довгий наступний абзац, який описує історію бренду та "
            "містить звичайний текст для публікації на сайті."
        )

    result = parse_docx(_docx_bytes(build), "article.docx")

    assert result.blocks[0].role == "h2"
    assert result.blocks[0].confidence >= 0.8


def test_long_single_paragraph_is_flagged_as_unstructured_canvas():
    def build(document):
        document.add_paragraph("Це великий нерозмічений текст. " * 20)

    result = parse_docx(_docx_bytes(build), "canvas.docx")

    assert [block.role for block in result.blocks] == ["p"]
    assert any("суцільним полотном" in warning for warning in result.warnings)


def test_long_document_without_headings_is_flagged_for_review():
    def build(document):
        document.add_paragraph("Перший звичайний абзац тексту. " * 10)
        document.add_paragraph("Другий звичайний абзац тексту. " * 10)

    result = parse_docx(_docx_bytes(build), "no-headings.docx")

    assert all(block.role == "p" for block in result.blocks)
    assert any("Заголовки не розпізнані" in warning for warning in result.warnings)


def test_short_document_without_headings_is_not_flagged():
    def build(document):
        document.add_paragraph("Короткий текст без заголовка.")

    result = parse_docx(_docx_bytes(build), "short.docx")

    assert not any("полотном" in warning for warning in result.warnings)
    assert not any("Заголовки не розпізнані" in warning for warning in result.warnings)


def test_pasted_text_detects_heading_paragraphs_and_lists():
    text = """Як обрати матеріал
Це достатньо довгий абзац, щоб короткий попередній рядок був визначений як заголовок для статті.
• Перший пункт
• Другий пункт"""

    result = parse_plain_text(text)

    assert [block.role for block in result.blocks] == ["h2", "p", "ul", "ul"]
    assert result.blocks[0].confidence == 0.55
    assert "<ul>" in render_html(result.blocks)
    assert "•" not in render_html(result.blocks)


def test_pasted_html_is_escaped_instead_of_executed():
    result = parse_plain_text("<script>alert('x')</script>")

    assert render_html(result.blocks) == (
        "<p>&lt;script&gt;alert('x')&lt;/script&gt;</p>"
    )


def test_rich_text_preserves_heading_ordered_list_and_link():
    result = parse_rich_text(
        "<h1>Як носити панаму?</h1>"
        "<p>Перевага сучасної панами — її універсальність.</p>"
        "<ol><li>біла сорочка</li><li>лляний костюм</li></ol>"
        '<p><a href="https://example.com/catalog">Перейти до каталогу</a></p>'
    )

    assert [block.role for block in result.blocks] == ["h2", "p", "ol", "ol", "p"]
    fragment = render_html(result.blocks)
    assert "<h2>Як носити панаму?</h2>" in fragment
    assert "<ol>" in fragment
    assert '<a href="https://example.com/catalog">' in fragment


def test_rich_text_uses_full_bold_short_paragraph_as_heading_signal():
    result = parse_rich_text(
        "<p><strong>Найпростіші поєднання</strong></p>"
        "<p>Це достатньо довгий наступний абзац для перевірки того, що попередній жирний рядок визначається як заголовок.</p>"
    )

    assert result.blocks[0].role == "h2"
    assert "<strong" not in render_html(result.blocks)


def test_google_docs_links_are_extracted_from_markdown_and_deduplicated():
    first = "https://docs.google.com/document/d/abc_123/edit?tab=t.0"
    second = "https://docs.google.com/document/u/0/d/xyz-789/edit"
    text = f"| [{first}]({first}) |\n| {second} |"

    assert extract_google_doc_links(text) == [first, second]
    assert extract_google_doc_id(first) == "abc_123"
    assert extract_google_doc_id(second) == "xyz-789"


def test_google_doc_download_checks_type_size_and_filename():
    response = Mock()
    response.status_code = 200
    response.headers = {
        "content-type": (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        "content-disposition": "attachment; filename*=UTF-8''Test%20Article.docx",
        "content-length": "6",
    }
    response.iter_content.return_value = [b"doc", b"x!!"]
    response.__enter__ = Mock(return_value=response)
    response.__exit__ = Mock(return_value=False)

    with patch("content_publisher.google_docs.requests.get", return_value=response):
        data, filename = download_google_doc(
            "https://docs.google.com/document/d/abc_123/edit", 10
        )

    assert data == b"docx!!"
    assert filename == "Test Article.docx"


def test_google_doc_download_rejects_non_docx_response():
    response = Mock()
    response.status_code = 200
    response.headers = {"content-type": "text/html"}
    response.__enter__ = Mock(return_value=response)
    response.__exit__ = Mock(return_value=False)

    with patch("content_publisher.google_docs.requests.get", return_value=response):
        try:
            download_google_doc("https://docs.google.com/document/d/abc/edit", 100)
        except GoogleDocsError as exc:
            assert "не віддав DOCX" in str(exc)
        else:
            raise AssertionError("Expected GoogleDocsError")


def test_leroy_merlin_renderer_uses_required_markup():
    blocks = [
        Block(
            index=1,
            text="Оберіть ПВХ-панелі для ремонту.",
            parts=[
                InlinePart("Оберіть "),
                InlinePart("ПВХ-панелі", "https://www.leroymerlin.ua/paneli"),
                InlinePart(" для ремонту."),
            ],
            role="p",
        ),
        Block(
            index=2,
            text="Рулетка для точного розкрою",
            parts=[InlinePart("Рулетка для точного розкрою")],
            role="ul",
        ),
    ]

    fragment = render_html(blocks, "leroy_merlin")

    assert 'data-gjs-tagName="a" data-gjs-type="text"' in fragment
    assert '<font color="#78BE20">ПВХ-панелі</font>' in fragment
    assert fragment.splitlines()[0].endswith("<br><br></p>")
    assert '<li data-gjs-tagName="li" data-gjs-type="text">' in fragment
    assert validate_html(fragment, "leroy_merlin") == []


def test_default_renderer_never_receives_leroy_markup():
    block = Block(
        index=1,
        text="Посилання",
        parts=[InlinePart("Посилання", "https://example.com")],
        role="p",
    )

    fragment = render_html([block], "default")

    assert fragment == '<p><a href="https://example.com">Посилання</a></p>'
    assert validate_html(fragment, "default") == []


def test_hyperlink_is_extracted_from_docx_and_rendered():
    def build(document):
        paragraph = document.add_paragraph("Перейдіть до ")
        _add_hyperlink(paragraph, "каталогу", "https://example.com/catalog")
        paragraph.add_run(".")

    result = parse_docx(_docx_bytes(build), "links.docx")
    fragment = render_html(result.blocks)

    assert fragment == (
        '<p>Перейдіть до <a href="https://example.com/catalog">каталогу</a>.</p>'
    )


def test_docx_export_contains_visible_html_source():
    fragment = "<h2>Заголовок</h2>\n<p>Текст без <strong>жирного</strong>.</p>"

    exported = Document(io.BytesIO(build_docx(fragment)))

    assert exported.paragraphs[0].text == fragment


def test_batch_export_contains_docx_files_with_html_source():
    result = DocumentResult(
        filename="Стаття.docx",
        blocks=[
            Block(
                index=1,
                text="Заголовок",
                parts=[InlinePart("Заголовок")],
                role="h2",
            )
        ],
    )

    archive = zipfile.ZipFile(io.BytesIO(build_docx_zip([result])))

    assert archive.namelist() == ["Стаття_HTML.docx"]
    exported = Document(io.BytesIO(archive.read("Стаття_HTML.docx")))
    assert exported.paragraphs[0].text == "<h2>Заголовок</h2>"
